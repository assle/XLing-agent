import smtplib
import ssl
import threading
from email.message import EmailMessage
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core import diagnostics
from app.core.config import Settings
from app.core.enums import ToolStatus
from app.models.entities import AlertRecord, ExcelRecord, SafetyAssessmentRecord, UserAccount

EXCEL_WRITE_LOCK = threading.RLock()


class ToolOrchestrationService:
    def __init__(self, db: Session, settings: Settings):
        """保存表格和通知工具使用的数据库会话及配置。

        初始化不创建文件或连接邮件服务器。
        """
        self.db = db
        self.settings = settings

    def write_excel(self, report: SafetyAssessmentRecord) -> ExcelRecord:
        """把安全评估记录追加到表格并登记写入结果。

        在进程内文件锁中确认评估仍存在；若已有成功记录则复用，否则打开或创建工作簿并保存。
        文件保存先于数据库提交，两者不共享事务，不能保证异常情况下绝不重复写行。
        """
        with diagnostics.stage("tool.excel", report_id=report.id):
            # 账户删除也持有此锁直到业务提交，等待中的工具不能重新导出已删除评估。
            with EXCEL_WRITE_LOCK:
                if self.db.query(SafetyAssessmentRecord.id).filter(SafetyAssessmentRecord.id == report.id).first() is None:
                    raise ValueError(f"report {report.id} not found")
                existing = (
                    self.db.query(ExcelRecord)
                    .filter(ExcelRecord.report_id == report.id, ExcelRecord.status == ToolStatus.SUCCESS.value)
                    .first()
                )
                # 复用数据库中的成功结果，避免普通重复请求再次执行外部操作。
                if existing is not None:
                    diagnostics.emit("tool.reused", tool="excel", report_id=report.id)
                    return existing
                path = Path(self.settings.excel_path)
                path.parent.mkdir(parents=True, exist_ok=True)
                if path.exists():
                    workbook = load_workbook(path)
                    sheet = workbook.active
                else:
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet.title = "Xling Risk Ledger"
                    sheet.append(["reportId", "riskLevel", "emotion", "confidence", "summary", "createdAt"])
                sheet.append([report.id, report.risk_level, report.emotion, report.confidence, report.summary, report.created_at.isoformat()])
                # 先把工作簿写到文件系统，下面才登记数据库结果；异常时两者可能处于不同状态。
                workbook.save(path)
                workbook.close()
                record = ExcelRecord(report_id=report.id, file_path=str(path), status=ToolStatus.SUCCESS.value, message="Excel 台账已写入")
                self.db.add(record)
                self.db.commit()
                return record

    @staticmethod
    def delete_excel_reports(report_ids: list[int], paths: list[str]) -> int:
        """从已知台账中移除指定评估行，保留表头和其他用户的记录。

        和写入共用进程内锁；文件缺失无需操作，读取或保存失败直接抛出。
        """
        if not report_ids:
            return 0
        deleted = 0
        with EXCEL_WRITE_LOCK:
            for path in {Path(value) for value in paths}:
                if not path.exists():
                    continue
                workbook = load_workbook(path)
                try:
                    sheet = workbook.active
                    if sheet.cell(1, 1).value != "reportId":
                        raise ValueError("Unrecognized risk ledger header")
                    rows = [
                        row for row in range(2, sheet.max_row + 1)
                        if sheet.cell(row, 1).value in report_ids
                    ]
                    for row in reversed(rows):
                        sheet.delete_rows(row)
                    if rows:
                        with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temporary:
                            temporary_path = Path(temporary.name)
                        try:
                            workbook.save(temporary_path)
                            temporary_path.replace(path)
                        finally:
                            temporary_path.unlink(missing_ok=True)
                    deleted += len(rows)
                finally:
                    workbook.close()
        return deleted

    def notify(self, report: SafetyAssessmentRecord, *, escalation: bool = False) -> AlertRecord:
        """按普通预警或超时升级类型投递通知，并分别记录真实结果。"""
        kind = "ESCALATION_ALERT" if escalation else "RISK_ALERT"
        label = "人工支持超时升级" if escalation else "高风险预警"
        with diagnostics.stage("tool.notification", report_id=report.id, notification_kind=kind):
            existing = (
                self.db.query(AlertRecord)
                .filter(
                    AlertRecord.report_id == report.id,
                    AlertRecord.notification_kind == kind,
                    AlertRecord.status == ToolStatus.SUCCESS.value,
                )
                .first()
            )
            if existing is not None:
                diagnostics.emit("tool.reused", tool="notification", report_id=report.id)
                return existing
            configured = (
                self.settings.escalation_alert_email_to.strip() or self.settings.alert_email_to.strip()
                if escalation else self.settings.alert_email_to.strip()
            )
            recipient = configured or "unconfigured"
            mode = self.settings.alert_email_delivery_mode.strip().lower()
            if mode == "log":
                return self._save_alert(
                    report, recipient if recipient != "unconfigured" else "log", ToolStatus.SUCCESS.value,
                    f"{label}已记录：reportId={report.id}，deliveryMode=log", notification_kind=kind,
                )
            if mode != "smtp":
                diagnostics.degraded("tool.notification", "unknown_delivery_mode", report_id=report.id)
                return self._save_alert(
                    report, recipient, ToolStatus.FAILED.value,
                    f"{label}邮件未发送：未知投递模式 {self.settings.alert_email_delivery_mode}",
                    notification_kind=kind,
                )
            missing = self._missing_email_config(escalation=escalation)
            if missing:
                diagnostics.degraded("tool.notification", "missing_email_configuration", report_id=report.id)
                return self._save_alert(
                    report, recipient, ToolStatus.FAILED.value,
                    f"{label}邮件未发送：缺少配置 {', '.join(missing)}", notification_kind=kind,
                )
            try:
                self._send_alert_email(report, escalation=escalation)
            except Exception as exc:
                diagnostics.degraded("tool.notification", "email_delivery_failed", exc, report_id=report.id)
                return self._save_alert(
                    report, recipient, ToolStatus.FAILED.value,
                    f"{label}邮件发送失败：{type(exc).__name__}: {exc}", notification_kind=kind,
                )
            return self._save_alert(
                report, recipient, ToolStatus.SUCCESS.value,
                f"{label}邮件已发送：reportId={report.id}", notification_kind=kind,
            )

    def _save_alert(
        self, report: SafetyAssessmentRecord, recipient: str, status: str, message: str,
        *, notification_kind: str = "RISK_ALERT",
    ) -> AlertRecord:
        """保存通知的收件人、状态和结果文字。"""
        record = AlertRecord(
            report_id=report.id, channel="email", notification_kind=notification_kind,
            recipient=recipient, status=status, message=message,
        )
        self.db.add(record)
        self.db.commit()
        return record

    def _missing_email_config(self, *, escalation: bool = False) -> list[str]:
        """检查邮件发送所需的服务器、发件人和收件人配置。"""
        missing = []
        if not self.settings.smtp_host.strip():
            missing.append("SMTP_HOST")
        if not self._sender():
            missing.append("ALERT_EMAIL_FROM 或 SMTP_USERNAME")
        if not self._recipients(escalation=escalation):
            missing.append("ESCALATION_ALERT_EMAIL_TO 或 ALERT_EMAIL_TO" if escalation else "ALERT_EMAIL_TO")
        return missing

    def _send_alert_email(self, report: SafetyAssessmentRecord, *, escalation: bool = False) -> None:
        """组装风险或升级通知并通过配置的邮件连接发送。"""
        message = EmailMessage()
        label = "人工支持超时升级" if escalation else self.settings.alert_email_subject_prefix
        message["Subject"] = f"{label} reportId={report.id}"
        message["From"] = self._sender()
        message["To"] = ", ".join(self._recipients(escalation=escalation))
        message.set_content(self._email_body(report))

        context = ssl.create_default_context()
        if self.settings.smtp_use_ssl:
            with smtplib.SMTP_SSL(
                self.settings.smtp_host,
                self.settings.smtp_port,
                timeout=self.settings.smtp_timeout_seconds,
                context=context,
            ) as server:
                self._send_message(server, message)
            return

        with smtplib.SMTP(self.settings.smtp_host, self.settings.smtp_port, timeout=self.settings.smtp_timeout_seconds) as server:
            server.ehlo()
            if self.settings.smtp_use_tls:
                server.starttls(context=context)
                server.ehlo()
            self._send_message(server, message)

    def _send_message(self, server: smtplib.SMTP, message: EmailMessage) -> None:
        """在已有邮件连接上按需登录并提交邮件。

        只有配置用户名时才登录；服务器接受邮件不等于收件人已经阅读或最终投递成功。
        """
        if self.settings.smtp_username:
            server.login(self.settings.smtp_username, self.settings.smtp_password)
        server.send_message(message)

    def _email_body(self, report: SafetyAssessmentRecord) -> str:
        """根据安全评估记录和关联用户信息构造邮件正文。

        当前正文包含用户标识、摘要及原始消息；不调用脱敏器，不应把它描述为脱敏摘要。
        """
        user = self.db.get(UserAccount, report.user_id)
        username = user.username if user else f"userId={report.user_id}"
        display_name = user.display_name if user else ""
        return "\n".join(
            [
                "Xling 检测到一条高风险安全预警，请尽快由授权审核团队安排跟进。",
                "",
                f"报告ID：{report.id}",
                f"用户：{display_name} ({username})" if display_name else f"用户：{username}",
                f"风险等级：{report.risk_level}",
                f"情绪标签：{report.emotion}",
                f"置信度：{report.confidence}",
                f"摘要：{report.summary}",
                f"创建时间：{report.created_at.isoformat()}",
                "",
                "用户原始消息：",
                report.content,
            ]
        )

    def _sender(self) -> str:
        """选择通知发件人地址。

        独立发件人配置优先，否则使用邮件登录用户名，并去除两端空白。
        """
        return self.settings.alert_email_from.strip() or self.settings.smtp_username.strip()

    def _recipients(self, *, escalation: bool = False) -> list[str]:
        """读取普通或备用升级通知接收方，不回退虚构地址。"""
        configured = (
            self.settings.escalation_alert_email_to or self.settings.alert_email_to
            if escalation else self.settings.alert_email_to
        )
        normalized = configured.replace(";", ",")
        return [recipient.strip() for recipient in normalized.split(",") if recipient.strip()]
