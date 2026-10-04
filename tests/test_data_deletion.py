"""Tests for issue 13: Privacy notice + user data deletion.

Covers:
  - Privacy notice accessible without login
  - Data deletion removes all user data
  - Deletion is transactional (rollback on failure)
  - After deletion, user account and JWT are invalid
  - Data deletion requires authentication

Run: python -m pytest tests/test_data_deletion.py
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.api.routes import router
from app.core.config import Settings
from app.core.security import hash_password
from app.models.entities import (
    ActionPlan,
    ActionPlanItem,
    ChatMessage,
    ChatSession,
    CheckIn,
    CheckpointDeletionTask,
    ExcelRecord,
    MemoryCard,
    RiskTrajectoryPoint,
    SafetyAssessmentRecord,
    ScreeningResult,
    UserAccount,
    UserProfile,
)
from app.services.action_plan import ActionPlanService
from app.services.data_deletion import DataDeletionService
from app.services.memory import RedisShortTermMemoryStore
from app.services.memory_cards import MemoryCardService
from app.services.tools import ToolOrchestrationService
from app.services.user_profile import UserProfileService
from tests.support import ApiHarness

_harness = ApiHarness(router)
_TestSession = _harness.sessions
client = _harness.client


class _Redis:
    def __init__(self):
        self.values = {}

    def rpush(self, key, value):
        self.values.setdefault(key, []).append(value)

    def ltrim(self, key, start, end):
        self.values[key] = self.values[key][start:]

    def expire(self, key, seconds):
        return True

    def set(self, key, value, ex=None):
        self.values[key] = value

    def delete(self, *keys):
        deleted = sum(key in self.values for key in keys)
        for key in keys:
            self.values.pop(key, None)
        return deleted


@pytest.fixture(autouse=True)
def isolated_external_data(tmp_path, monkeypatch):
    settings = Settings(
        _env_file=None,
        excel_path=str(tmp_path / "ledger.xlsx"),
        langgraph_checkpoint_backend="memory",
    )
    cache = _Redis()
    monkeypatch.setattr("app.services.data_deletion.get_settings", lambda: settings)
    monkeypatch.setattr(RedisShortTermMemoryStore, "_connect", lambda self: cache)
    monkeypatch.setattr(RedisShortTermMemoryStore, "_fallback", {})
    yield settings, cache

def _seed():
    """创建两个普通用户用于数据删除隔离测试。

    让被删除用户与应保留用户有独立账户。
    """
    db = _TestSession()
    try:
        s = UserAccount(username="student", display_name="S", password_hash=hash_password("student123"))
        s.roles = {"ROLE_USER"}
        s2 = UserAccount(username="student2", display_name="S2", password_hash=hash_password("p2"))
        s2.roles = {"ROLE_USER"}
        db.add_all([s, s2])
        db.commit()
    finally:
        db.close()

_seed()

def _token(u="student", p="student123"):
    """通过指定测试账户登录取得凭证。

    可在删除前后用同一凭证检查账户是否仍有效。
    """
    r = client.post("/api/auth/login", json={"username": u, "password": p})
    return r.json()["accessToken"]

def _auth(t):
    """构造携带指定凭证的请求头。

    保留凭证原文，便于删除后再次尝试访问。
    """
    return {"Authorization": f"Bearer {t}"}

def _setup_user_data(user_id=1):
    """为用户建立背景、卡片、筛查、计划、轨迹和会话评估数据。

    模拟跨多张表的关联数据，供删除数量与隔离断言使用。
    """
    db = _TestSession()
    try:
        # Profile
        UserProfileService(db).update_support_background(user_id, current_concern="近期压力较大")
        # Memory cards
        MemoryCardService(db).create_card(user_id, "important memory")
        # Screening result
        from app.services.screening import ScreeningService
        ScreeningService(db).submit_screening(user_id, "PHQ-9", [0]*9)
        # Action plan
        ActionPlanService(db, ai=None).generate_plan(user_id, None, "summary")
        # Risk trajectory
        from app.core.enums import RiskLevel
        from app.services.risk_trajectory import RiskTrajectoryService
        RiskTrajectoryService(db).record_point(user_id, None, RiskLevel.LOW, 1.0)
        # Chat session + message
        session = ChatSession(public_id=f"test-sess-{user_id}", title="test", user_id=user_id)
        db.add(session)
        db.flush()
        db.add(ChatMessage(user_id=user_id, session_id=session.id, role="USER", content="test"))
        db.add(SafetyAssessmentRecord(
            user_id=user_id, session_id=session.id, content="test",
            intent="CONSULT", emotion="ANXIETY", emotion_score=2.0,
            risk_level="LOW", confidence=0.7, summary="test",
        ))
        db.commit()
    finally:
        db.close()

def _reset_db():
    """按依赖顺序清理本文件测试数据，再重建两个账户。

    此文件存在同名重复定义，注释覆盖两处，运行时后一个定义生效。
    """
    db = _TestSession()
    try:
        for model in [ExcelRecord, CheckpointDeletionTask, CheckIn, ActionPlanItem, ActionPlan, RiskTrajectoryPoint, ChatMessage,
                      SafetyAssessmentRecord, ChatSession, ScreeningResult, MemoryCard, UserProfile, UserAccount]:
            db.query(model).delete()
        db.commit()
    finally:
        db.close()
    _seed()


def _reset_db():
    """按依赖顺序清理本文件测试数据，再重建两个账户。

    此文件存在同名重复定义，注释覆盖两处，运行时后一个定义生效。
    """
    db = _TestSession()
    try:
        for model in [ExcelRecord, CheckpointDeletionTask, CheckIn, ActionPlanItem, ActionPlan, RiskTrajectoryPoint, ChatMessage,
                      SafetyAssessmentRecord, ChatSession, ScreeningResult, MemoryCard, UserProfile, UserAccount]:
            db.query(model).delete()
        db.commit()
    finally:
        db.close()
    _seed()


# ---------------------------------------------------------------------------
# Privacy notice
# ---------------------------------------------------------------------------

def test_privacy_notice_no_auth_required():
    """匿名访问隐私说明接口。

    检查无需登录且正文包含收集、删除和无记忆相关说明。
    """
    response = client.get("/api/privacy")
    assert response.status_code == 200
    data = response.json()
    assert "notice" in data
    assert "隐私" in data["notice"]
    assert "收集" in data["notice"]
    assert "删除" in data["notice"]
    assert "无记忆" in data["notice"]


# ---------------------------------------------------------------------------
# Data deletion
# ---------------------------------------------------------------------------

def test_delete_removes_all_user_data():
    """为用户准备多类业务数据后调用删除服务。

    核对各类删除数量及账户不存在，验证当前测试列出的关联记录。
    """
    _reset_db()
    _setup_user_data(1)
    db = _TestSession()
    try:
        svc = DataDeletionService(db)
        counts = svc.delete_all_user_data(1)
        assert counts["user_account"] == 1
        assert counts["memory_cards"] >= 1
        assert counts["screening_results"] >= 1
        assert counts["user_profiles"] >= 1
        assert counts["action_plans"] >= 1
        assert counts["risk_trajectory"] >= 1
        assert counts["chat_messages"] >= 1
        assert counts["safety_assessment_records"] >= 1
        # User is gone
        assert db.get(UserAccount, 1) is None
    finally:
        db.close()

def test_delete_does_not_affect_other_users():
    """分别建立两位用户的数据，只删除第一位。

    检查第二位账户和支持背景仍存在。
    """
    _reset_db()
    _setup_user_data(1)
    _setup_user_data(2)
    db = _TestSession()
    try:
        svc = DataDeletionService(db)
        svc.delete_all_user_data(1)
        # User 2 still exists
        assert db.get(UserAccount, 2) is not None
        assert UserProfileService(db).get_profile(2) is not None
    finally:
        db.close()

def test_delete_makes_token_invalid():
    """先验证凭证可访问资料，再用它删除账户并再次访问。

    检查账户删除后原凭证无法继续通过当前用户查询。
    """
    _reset_db()
    _setup_user_data(1)
    token = _token()
    # Verify token works
    r = client.get("/api/profile", headers=_auth(token))
    assert r.status_code == 200
    # Delete account
    r = client.delete("/api/account", headers=_auth(token))
    assert r.status_code == 200
    # Token should now be invalid (user not found)
    r = client.get("/api/profile", headers=_auth(token))
    assert r.status_code == 401

def test_delete_requires_auth():
    """未携带凭证请求删除账户。

    检查接口拒绝未经身份验证的操作。
    """
    r = client.delete("/api/account")
    assert r.status_code == 401

def test_api_delete_account():
    """通过接口删除已准备数据的测试用户。

    检查响应报告删除完成并包含详情。
    """
    _reset_db()
    _setup_user_data(1)
    token = _token()
    r = client.delete("/api/account", headers=_auth(token))
    assert r.status_code == 200
    data = r.json()
    assert data["deleted"] is True
    assert "details" in data


def test_data_deletion_removes_persistent_checkpoints(tmp_path):
    """在临时状态数据库中创建该用户的快照和写入行。

    删除业务数据后再读文件，检查两类执行状态记录都已清理。
    """
    _reset_db()
    _setup_user_data(1)
    db = _TestSession()
    checkpoint_path = tmp_path / "checkpoints.db"
    try:
        thread_id = db.query(ChatSession).filter(ChatSession.user_id == 1).one().public_id
        connection = sqlite3.connect(checkpoint_path)
        connection.executescript(
            """
            CREATE TABLE checkpoints (
                thread_id TEXT, checkpoint_ns TEXT, checkpoint_id TEXT,
                parent_checkpoint_id TEXT, type TEXT, checkpoint BLOB, metadata BLOB
            );
            CREATE TABLE writes (
                thread_id TEXT, checkpoint_ns TEXT, checkpoint_id TEXT,
                task_id TEXT, idx INTEGER, channel TEXT, type TEXT, value BLOB
            );
            """
        )
        connection.execute(
            "INSERT INTO checkpoints(thread_id, checkpoint_id) VALUES (?, ?)",
            (thread_id, "checkpoint-1"),
        )
        connection.execute(
            "INSERT INTO writes(thread_id, checkpoint_id) VALUES (?, ?)",
            (thread_id, "checkpoint-1"),
        )
        connection.commit()
        connection.close()

        DataDeletionService(
            db,
            Settings(
                _env_file=None,
                excel_path=str(tmp_path / "ledger.xlsx"),
                langgraph_checkpoint_backend="async_sqlite",
                langgraph_checkpoint_path=str(checkpoint_path),
            ),
        ).delete_all_user_data(1)

        connection = sqlite3.connect(checkpoint_path)
        assert connection.execute("SELECT COUNT(*) FROM checkpoints").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM writes").fetchone()[0] == 0
        connection.close()
    finally:
        db.close()


def test_business_rollback_does_not_delete_checkpoints(tmp_path, monkeypatch):
    """让业务提交主动失败。

    检查异常向外传递、账户仍存在且文件执行状态未被提前删除。
    """
    _reset_db()
    _setup_user_data(1)
    db = _TestSession()
    checkpoint_path = tmp_path / "checkpoints.db"
    try:
        thread_id = (
            db.query(ChatSession)
            .filter(ChatSession.user_id == 1)
            .one()
            .public_id
        )
        connection = sqlite3.connect(checkpoint_path)
        connection.execute(
            "CREATE TABLE checkpoints (thread_id TEXT, checkpoint_id TEXT)"
        )
        connection.execute(
            "INSERT INTO checkpoints(thread_id, checkpoint_id) VALUES (?, ?)",
            (thread_id, "checkpoint-1"),
        )
        connection.commit()
        connection.close()

        def fail_commit():
            """模拟业务数据库确认保存失败。

            迫使删除流程回滚，供测试核对外部状态仍保留。
            """
            raise RuntimeError("business commit failed")

        monkeypatch.setattr(db, "commit", fail_commit)

        with pytest.raises(RuntimeError, match="business commit failed"):
            DataDeletionService(
                db,
                Settings(
                    _env_file=None,
                    excel_path=str(tmp_path / "ledger.xlsx"),
                    langgraph_checkpoint_backend="async_sqlite",
                    langgraph_checkpoint_path=str(checkpoint_path),
                ),
            ).delete_all_user_data(1)

        connection = sqlite3.connect(checkpoint_path)
        assert connection.execute("SELECT COUNT(*) FROM checkpoints").fetchone()[0] == 1
        connection.close()
        assert db.get(UserAccount, 1) is not None
    finally:
        db.close()


def test_failed_checkpoint_cleanup_is_persisted_and_retried(tmp_path, monkeypatch):
    """让业务删除成功但文件状态清理失败。

    检查账户已删除、待清理任务与次数保留；恢复清理函数后重试应清除任务和状态行。
    """
    _reset_db()
    _setup_user_data(1)
    db = _TestSession()
    checkpoint_path = tmp_path / "checkpoints.db"
    try:
        thread_id = (
            db.query(ChatSession)
            .filter(ChatSession.user_id == 1)
            .one()
            .public_id
        )
        connection = sqlite3.connect(checkpoint_path)
        connection.execute(
            "CREATE TABLE checkpoints (thread_id TEXT, checkpoint_id TEXT)"
        )
        connection.execute(
            "INSERT INTO checkpoints(thread_id, checkpoint_id) VALUES (?, ?)",
            (thread_id, "checkpoint-1"),
        )
        connection.commit()
        connection.close()
        service = DataDeletionService(
            db,
            Settings(
                _env_file=None,
                excel_path=str(tmp_path / "ledger.xlsx"),
                langgraph_checkpoint_backend="async_sqlite",
                langgraph_checkpoint_path=str(checkpoint_path),
            ),
        )
        original_delete = service._delete_checkpoint_rows

        def fail_cleanup(thread_ids, path):
            """模拟独立状态数据库被占用而无法删除。

            让业务删除后的补偿任务保留到下次重试。
            """
            raise sqlite3.OperationalError("database busy")

        monkeypatch.setattr(service, "_delete_checkpoint_rows", fail_cleanup)

        counts = service.delete_all_user_data(1)

        assert counts["checkpoint_cleanup_pending"] == 1
        task = db.query(CheckpointDeletionTask).one()
        assert task.attempts == 1
        assert db.get(UserAccount, 1) is None

        monkeypatch.setattr(service, "_delete_checkpoint_rows", original_delete)
        assert service.retry_pending_checkpoint_deletions() == 1
        assert db.query(CheckpointDeletionTask).count() == 0
        connection = sqlite3.connect(checkpoint_path)
        assert connection.execute("SELECT COUNT(*) FROM checkpoints").fetchone()[0] == 0
        connection.close()
    finally:
        db.close()


def test_delete_clears_owned_cache_and_all_known_ledger_paths(isolated_external_data, tmp_path):
    """删除一个用户的外部原文和导出，同时保留另一个用户的数据。"""
    _reset_db()
    _setup_user_data(1)
    _setup_user_data(2)
    settings, cache = isolated_external_data
    db = _TestSession()
    try:
        memory = RedisShortTermMemoryStore(settings)
        for user_id in (1, 2):
            thread_id = f"test-sess-{user_id}"
            memory.append(thread_id, "USER", f"private message {user_id}")
            memory.save_cbt_state(thread_id, {"event": f"private event {user_id}"})
            RedisShortTermMemoryStore._fallback[memory._key(thread_id)] = [f"fallback {user_id}"]
            RedisShortTermMemoryStore._fallback[memory._cbt_key(thread_id)] = [f"CBT fallback {user_id}"]
        reports = db.query(SafetyAssessmentRecord).order_by(SafetyAssessmentRecord.user_id).all()
        tools = ToolOrchestrationService(db, settings)
        for report in reports:
            tools.write_excel(report)
        report_ids = [report.id for report in reports]
        historical_path = tmp_path / "previous-ledger.xlsx"
        workbook = load_workbook(settings.excel_path)
        retained_row = tuple(workbook.active.values)[2]
        # 失败后的重复写行也属于被删除用户的数据。
        workbook.active.append([report_ids[0], "LOW", "duplicate"])
        workbook.save(settings.excel_path)
        workbook.save(historical_path)
        workbook.close()
        db.add(ExcelRecord(
            report_id=report_ids[0], file_path=str(historical_path),
            status="SUCCESS", message="previous export",
        ))
        db.commit()
        other_cache = {key: value for key, value in cache.values.items() if key.endswith("test-sess-2")}

        counts = DataDeletionService(db, settings).delete_all_user_data(1)

        assert counts["cached_session_keys"] == 2
        assert counts["ledger_rows"] == 4
        assert cache.values == other_cache
        assert set(RedisShortTermMemoryStore._fallback) == {
            memory._key("test-sess-2"), memory._cbt_key("test-sess-2"),
        }
        for path in (settings.excel_path, historical_path):
            workbook = load_workbook(path)
            assert list(workbook.active.values) == [
                ("reportId", "riskLevel", "emotion", "confidence", "summary", "createdAt"),
                retained_row,
            ]
            workbook.close()
        assert db.get(UserAccount, 2) is not None
        assert db.get(SafetyAssessmentRecord, report_ids[1]) is not None
    finally:
        db.close()


def test_cache_cleanup_failure_rolls_back_account_deletion(isolated_external_data, monkeypatch):
    _reset_db()
    _setup_user_data(1)
    settings, cache = isolated_external_data
    db = _TestSession()
    try:
        tools = ToolOrchestrationService(db, settings)
        report = db.query(SafetyAssessmentRecord).one()
        tools.write_excel(report)
        before = Path(settings.excel_path).read_bytes()

        def fail_delete(*keys):
            raise ConnectionError("cache unavailable")

        monkeypatch.setattr(cache, "delete", fail_delete)
        with pytest.raises(ConnectionError, match="cache unavailable"):
            DataDeletionService(db, settings).delete_all_user_data(1)

        assert db.get(UserAccount, 1) is not None
        assert db.query(ChatMessage).count() == 1
        assert db.query(CheckpointDeletionTask).count() == 0
        assert Path(settings.excel_path).read_bytes() == before
    finally:
        db.close()


def test_ledger_cleanup_failure_keeps_account_and_can_be_retried(isolated_external_data, monkeypatch):
    _reset_db()
    _setup_user_data(1)
    settings, cache = isolated_external_data
    db = _TestSession()
    try:
        report = db.query(SafetyAssessmentRecord).one()
        ToolOrchestrationService(db, settings).write_excel(report)
        original_delete = ToolOrchestrationService.delete_excel_reports

        def fail_delete(report_ids, paths):
            raise OSError("export is read only")

        monkeypatch.setattr(ToolOrchestrationService, "delete_excel_reports", fail_delete)
        response = client.delete("/api/account", headers=_auth(_token()))
        assert response.status_code == 500
        assert db.get(UserAccount, 1) is not None
        assert db.query(SafetyAssessmentRecord).count() == 1
        monkeypatch.setattr(ToolOrchestrationService, "delete_excel_reports", original_delete)

        response = client.delete("/api/account", headers=_auth(_token()))
        assert response.status_code == 200
        assert response.json()["deleted"] is True
        assert response.json()["details"]["ledger_rows"] == 1
    finally:
        db.close()


def test_stale_tool_writer_cannot_export_deleted_report(isolated_external_data):
    _reset_db()
    _setup_user_data(1)
    settings, _ = isolated_external_data
    stale_db = _TestSession()
    deletion_db = _TestSession()
    try:
        stale_report = stale_db.query(SafetyAssessmentRecord).one()
        DataDeletionService(deletion_db, settings).delete_all_user_data(1)

        with pytest.raises(ValueError, match="report .* not found"):
            ToolOrchestrationService(stale_db, settings).write_excel(stale_report)
        assert stale_db.query(ExcelRecord).count() == 0
        assert not Path(settings.excel_path).exists()
    finally:
        stale_db.close()
        deletion_db.close()


def test_partial_ledger_save_failure_preserves_original_and_account(isolated_external_data, monkeypatch):
    _reset_db()
    _setup_user_data(1)
    _setup_user_data(2)
    settings, _ = isolated_external_data
    db = _TestSession()
    try:
        for report in db.query(SafetyAssessmentRecord).all():
            ToolOrchestrationService(db, settings).write_excel(report)
        ledger = Path(settings.excel_path)
        original = ledger.read_bytes()

        def partial_save(workbook, destination):
            Path(destination).write_bytes(b"partial workbook")
            raise OSError("disk full")

        monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", partial_save)
        with pytest.raises(OSError, match="disk full"):
            DataDeletionService(db, settings).delete_all_user_data(1)

        assert ledger.read_bytes() == original
        assert list(ledger.parent.iterdir()) == [ledger]
        assert db.get(UserAccount, 1) is not None
        assert db.get(UserAccount, 2) is not None
        assert db.query(SafetyAssessmentRecord).count() == 2
    finally:
        db.close()
